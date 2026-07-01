from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi.testclient import TestClient

from store_ai_clinic.api.main import app
from store_ai_clinic.api.routers import quality as quality_router
from store_ai_clinic.config import settings
from store_ai_clinic.services.quality_vision import VisionQualityFinding
from store_ai_clinic.services.quality_assets import extract_quality_asset_summary
from store_ai_clinic.services.quality_rules import extract_quality_rules
from store_ai_clinic.services.quality import (
    QualityDatasetScan as ServiceQualityDatasetScan,
    QualityExportSummary as ServiceQualityExportSummary,
    QualityIssue as ServiceQualityIssue,
    QualityMetric as ServiceQualityMetric,
    QualityPipelineStage as ServiceQualityPipelineStage,
    build_quality_export_summary,
    create_export_task,
    list_export_tasks,
    list_quality_actions,
    list_review_records,
    record_quality_action,
    record_review_decision,
    scan_quality_dataset,
)



@pytest.fixture(autouse=True)
def disable_default_quality_vision(monkeypatch):
    monkeypatch.setattr(settings, "openai_api_key", None, raising=False)
    monkeypatch.setattr(settings, "openai_vision_pdf_enabled", False, raising=False)

def write_pdf(path: Path, page_count: int) -> None:
    markers = "\n".join("/Type /Page" for _ in range(page_count))
    path.write_bytes(("%PDF-1.4\n" + markers + "\n%%EOF").encode("latin1"))


def write_xlsx(path: Path, rows: list[list[str]]) -> None:
    shared: list[str] = []

    def shared_index(value: str) -> int:
        shared.append(value)
        return len(shared) - 1

    row_xml = []
    for row_number, row in enumerate(rows, start=1):
        cells = []
        for column_index, value in enumerate(row):
            column = chr(ord("A") + column_index)
            cells.append(f'<c r="{column}{row_number}" t="s"><v>{shared_index(value)}</v></c>')
        row_xml.append(f'<row r="{row_number}">{"".join(cells)}</row>')

    shared_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{len(shared)}" uniqueCount="{len(shared)}">'
        + "".join(f"<si><t>{value}</t></si>" for value in shared)
        + "</sst>"
    )
    sheet_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        + "".join(row_xml)
        + "</sheetData></worksheet>"
    )

    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("xl/sharedStrings.xml", shared_xml)
        archive.writestr("xl/worksheets/sheet1.xml", sheet_xml)


def write_docx(path: Path, paragraphs: list[str]) -> None:
    paragraph_xml = "".join(f"<w:p><w:r><w:t>{paragraph}</w:t></w:r></w:p>" for paragraph in paragraphs)
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{paragraph_xml}</w:body>"
        "</w:document>"
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document_xml)
def build_quality_sample(root: Path) -> Path:
    dataset = root / "sample-dataset"
    group = dataset / "35-44"
    group.mkdir(parents=True)
    write_pdf(group / "02496166_report.pdf", page_count=3)
    write_xlsx(
        group / "structured-data.xlsx",
        [
            ["ArchivesNum"],
            ["02496166"],
        ],
    )
    (dataset / "requirements.docx").write_text("PDF + structured data", encoding="utf-8")
    return dataset


def test_quality_rules_are_extracted_from_requirement_docx(tmp_path: Path):
    dataset = tmp_path / "5人"
    dataset.mkdir()
    write_docx(
        tmp_path / "体检报告需求.docx",
        [
            "客户要求 PDF + 结构化数据。",
            "数据必须为 5 年内，同一个人连续 3 次就诊。",
            "关注页数分布、异常项数量、检查项目丰富度、历史数据对比、是否包含总结。",
        ],
    )

    rules = extract_quality_rules(dataset)

    rule_ids = {rule.rule_id for rule in rules}
    assert {
        "R-REQ-001",
        "R-REQ-002",
        "R-REQ-003",
        "R-REQ-004",
        "R-REQ-005",
        "R-REQ-006",
        "R-REQ-007",
        "R-REQ-008",
    }.issubset(rule_ids)
    summary_rule = next(rule for rule in rules if rule.rule_id == "R-REQ-008")
    assert summary_rule.source.endswith("体检报告需求.docx")
    assert summary_rule.dimension == "总结完整性"
    assert summary_rule.check_target == "PDF 报告"
    assert summary_rule.need_human_review is True


def test_quality_api_exposes_extracted_requirement_rules(tmp_path: Path, monkeypatch):
    dataset = tmp_path / "5人"
    dataset.mkdir()
    write_docx(
        tmp_path / "体检报告需求.docx",
        [
            "第三批样本需包含 PDF + 结构化数据。",
            "同一个人连续 3 次就诊，且检查日期在 5 年内。",
            "需要历史数据对比和体检总结。",
        ],
    )
    monkeypatch.setattr(settings, "local_storage_root", str(tmp_path / "api-state"))
    client = TestClient(app)

    response = client.get("/api/quality/rules", params={"dataset_path": str(dataset)})

    assert response.status_code == 200
    body = response.json()
    assert body["dataset_path"] == str(dataset)
    assert body["source_document"].endswith("体检报告需求.docx")
    assert {rule["rule_id"] for rule in body["rules"]} >= {"R-REQ-001", "R-REQ-002", "R-REQ-003", "R-REQ-007", "R-REQ-008"}



def test_quality_actions_are_recorded_and_queryable(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "local_storage_root", str(tmp_path / "api-state"))
    client = TestClient(app)

    response = client.post(
        "/api/quality/actions",
        json={
            "action": "open_report_detail",
            "label": "查看单报告详情",
            "page": "tasks",
            "target": "/agent",
            "dataset_path": str(tmp_path / "5人"),
            "payload": {"issue_id": "issue-1"},
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["action"] == "open_report_detail"
    assert body["label"] == "查看单报告详情"
    assert body["message"] == "已记录操作：查看单报告详情"

    records = client.get("/api/quality/actions", params={"dataset_path": str(tmp_path / "5人")})

    assert records.status_code == 200
    assert [item["action"] for item in records.json()] == ["open_report_detail"]


def test_quality_action_service_persists_to_jsonl(tmp_path: Path):
    record_quality_action(
        action="toggle_rule_status",
        label="切换 R-OCR-001 状态",
        page="rules",
        target="R-OCR-001",
        dataset_path=tmp_path / "5人",
        payload={"next_status": "启用"},
        storage_root=tmp_path / "state",
    )

    records = list_quality_actions(dataset_path=tmp_path / "5人", storage_root=tmp_path / "state")

    assert len(records) == 1
    assert records[0].action == "toggle_rule_status"
    assert records[0].payload == {"next_status": "启用"}

def test_quality_assets_map_archives_pdfs_excels_and_visit_counts(tmp_path: Path):
    dataset = tmp_path / "5人"
    group = dataset / "35-44"
    group.mkdir(parents=True)
    write_pdf(group / "02496166.pdf", page_count=6)
    write_pdf(group / "02496166-复查.pdf", page_count=5)
    write_pdf(group / "02497000.pdf", page_count=5)
    write_xlsx(
        group / "结构化数据.xlsx",
        [
            ["ArchivesNum", "CheckDate", "ItemFlag"],
            ["02496166", "2024-01-01", "血常规"],
            ["02496166", "2023-01-01", "血常规"],
            ["02496166", "2022-01-01", "血常规"],
            ["02498000", "2024-02-02", "尿常规"],
            ["02498000", "2023-02-02", "尿常规"],
        ],
    )

    summary = extract_quality_asset_summary(dataset, now="2026-06-30T12:00:00+08:00")

    assert summary.dataset_path == str(dataset)
    assert summary.total_groups == 1
    assert summary.total_archives == 3
    assert summary.total_pdf_files == 3
    assert summary.total_excel_files == 1
    assert summary.matched_archives == 1
    assert summary.missing_pdf_archives == 1
    assert summary.missing_excel_archives == 1
    assert summary.under_three_visit_archives == 2

    archive_with_match = next(asset for asset in summary.assets if asset.archive_id == "02496166")
    assert archive_with_match.group == "35-44"
    assert archive_with_match.pdf_files == ["02496166-复查.pdf", "02496166.pdf"]
    assert archive_with_match.excel_files == ["结构化数据.xlsx"]
    assert archive_with_match.visit_count == 3
    assert archive_with_match.has_pdf is True
    assert archive_with_match.has_excel is True
    assert archive_with_match.meets_three_visits is True

    missing_pdf = next(asset for asset in summary.assets if asset.archive_id == "02498000")
    assert missing_pdf.has_pdf is False
    assert missing_pdf.has_excel is True
    assert missing_pdf.missing_items == ["missing_pdf", "under_three_visits"]

    missing_excel = next(asset for asset in summary.assets if asset.archive_id == "02497000")
    assert missing_excel.has_pdf is True
    assert missing_excel.has_excel is False
    assert missing_excel.missing_items == ["missing_excel", "under_three_visits"]


def test_quality_api_exposes_asset_summary(tmp_path: Path):
    dataset = tmp_path / "5人"
    group = dataset / "45-54"
    group.mkdir(parents=True)
    write_pdf(group / "03110001.pdf", page_count=6)
    write_xlsx(
        group / "结构化数据.xlsx",
        [
            ["ArchivesNum", "CheckDate"],
            ["03110001", "2024-01-01"],
            ["03110001", "2023-01-01"],
            ["03110001", "2022-01-01"],
        ],
    )
    client = TestClient(app)

    response = client.get("/api/quality/assets", params={"dataset_path": str(dataset)})

    assert response.status_code == 200
    body = response.json()
    assert body["dataset_path"] == str(dataset)
    assert body["matched_archives"] == 1
    assert body["assets"][0]["archive_id"] == "03110001"
    assert body["assets"][0]["group"] == "45-54"

def test_quality_service_scans_dataset_and_preserves_evidence(tmp_path: Path):
    dataset = build_quality_sample(tmp_path)

    scan = scan_quality_dataset(dataset, now="2026-06-26T12:00:00+08:00")

    assert scan.dataset_path == str(dataset)
    assert len(scan.metrics) == 5
    assert len(scan.pipeline) == 5
    assert {issue.rule_id for issue in scan.issues} == {"R-FORMAT-002", "R-OCR-001", "R-EXCEL-001"}
    assert all(issue.evidence for issue in scan.issues)

    review = record_review_decision(
        issue_id=scan.issues[0].id,
        decision="confirmed",
        reviewer="tester",
        note="evidence accepted",
        evidence=scan.issues[0].evidence,
        storage_root=tmp_path / "service-state",
    )
    assert review.issue_id == scan.issues[0].id
    assert review.decision == "confirmed"

    export = create_export_task("review-records", dataset_path=dataset)
    assert export.export_type == "review-records"
    assert export.status == "done"
    assert export.bundle_name and export.bundle_name.endswith(".zip")
    assert export.download_url == f"/api/quality/exports/{export.id}/download"
    assert "已生成交付包" in export.message

def test_quality_api_exposes_scan_filter_review_and_export(tmp_path: Path, monkeypatch):
    dataset = build_quality_sample(tmp_path)
    monkeypatch.setattr(settings, "local_storage_root", str(tmp_path / "api-state"))
    client = TestClient(app)

    response = client.get("/api/quality/datasets/scan", params={"dataset_path": str(dataset)})
    assert response.status_code == 200
    body = response.json()
    assert body["dataset_path"] == str(dataset)
    assert len(body["issues"]) == 3

    filtered = client.get(
        "/api/quality/issues",
        params={"dataset_path": str(dataset), "severity": "high", "status": "needs_review"},
    )
    assert filtered.status_code == 200
    filtered_body = filtered.json()
    assert filtered_body["total"] >= 1
    assert filtered_body["issues"][0]["severity"] == "high"

    review = client.post(
        "/api/quality/reviews",
        json={
            "issue_id": filtered_body["issues"][0]["id"],
            "decision": "confirmed",
            "reviewer": "frontend-reviewer",
            "note": "confirmed by human reviewer",
            "evidence": filtered_body["issues"][0]["evidence"],
        },
    )
    assert review.status_code == 200
    assert review.json()["decision"] == "confirmed"

    export = client.post(
        "/api/quality/exports",
        json={"export_type": "review-records", "dataset_path": str(dataset)},
    )
    assert export.status_code == 200
    assert export.json()["status"] == "done"


def test_quality_api_imports_files_and_starts_scan_task(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "local_storage_root", str(tmp_path / "api-state"))
    client = TestClient(app)

    response = client.post(
        "/api/quality/import",
        files=[
            (
                "files",
                (
                    "../unsafe/report one.pdf",
                    b"%PDF-1.4\n/Type /Page\n/Type /Page\n%%EOF",
                    "application/pdf",
                ),
            )
        ],
    )

    assert response.status_code == 201
    body = response.json()
    assert body["data_source"] == "backend"
    assert body["persisted"] is True
    assert body["total_files"] == 1
    assert body["total_bytes"] > 0
    assert body["task_id"].startswith("quality-import-")
    dataset_path = Path(body["dataset_path"])
    assert dataset_path.exists()
    assert dataset_path.is_dir()
    assert body["files"][0]["name"] == "report one.pdf"
    assert (dataset_path / "report one.pdf").exists()

    scan_response = client.post(f"/api/quality/imports/{body['task_id']}/scan")

    assert scan_response.status_code == 200
    scan_body = scan_response.json()
    assert scan_body["task_id"] == body["task_id"]
    assert scan_body["status"] == "done"
    assert scan_body["scan"]["dataset_path"] == str(dataset_path)
    assert {issue["rule_id"] for issue in scan_body["scan"]["issues"]} >= {"R-FORMAT-002", "R-OCR-001"}


def test_quality_scan_uses_vision_model_findings_for_images(tmp_path: Path):
    dataset = tmp_path / "image-dataset"
    dataset.mkdir()
    image_path = dataset / "report-page.png"
    image_path.write_bytes(b"\x89PNG\r\n\x1a\nimage-bytes")
    captured: dict[str, Path] = {}

    class FakeVisionAnalyzer:
        def analyze_image(self, path: Path) -> VisionQualityFinding:
            captured["path"] = path
            return VisionQualityFinding(
                issue_type="AI 视觉判断",
                severity="high",
                evidence="图片识别到血红蛋白 88 g/L，低于参考范围 110-150 g/L。",
                ai_judgement="存在疑似贫血相关异常，需要人工复核报告原图。",
                recommendation="复核血常规页的血红蛋白项目，并关联结构化数据确认是否应标记异常。",
                confidence=0.88,
                finding_type="data_error",
                bbox=[12.0, 18.0, 36.0, 14.0],
            )

    scan = scan_quality_dataset(
        dataset,
        now="2026-06-26T12:00:00+08:00",
        vision_analyzer=FakeVisionAnalyzer(),
    )

    image_issues = [issue for issue in scan.issues if issue.file_name == "report-page.png"]
    assert captured["path"] == image_path
    assert len(image_issues) == 1
    assert image_issues[0].rule_id == "R-VISION-001"
    assert image_issues[0].issue_type == "AI 视觉判断"
    assert image_issues[0].severity == "high"
    assert image_issues[0].evidence == "图片识别到血红蛋白 88 g/L，低于参考范围 110-150 g/L。"
    assert image_issues[0].ai_judgement == "存在疑似贫血相关异常，需要人工复核报告原图。"
    assert image_issues[0].recommendation == "复核血常规页的血红蛋白项目，并关联结构化数据确认是否应标记异常。"
    assert image_issues[0].confidence == 0.88
    assert image_issues[0].finding_type == "data_error"
    assert image_issues[0].bbox == [12.0, 18.0, 36.0, 14.0]


def test_quality_scan_returns_pdf_page_preview_for_ocr_fallback(tmp_path: Path):
    dataset = tmp_path / "pdf-preview-dataset"
    dataset.mkdir()
    pdf_path = dataset / "02496056.pdf"
    write_pdf(pdf_path, page_count=6)
    rendered_path = tmp_path / "02496056-page-1.png"
    calls: dict[str, Path] = {}

    class FakePdfPageRenderer:
        def render_first_page(self, path: Path) -> Path:
            calls["pdf_path"] = path
            rendered_path.write_bytes(b"\x89PNG\r\n\x1a\nrendered-pdf-page")
            return rendered_path

    scan = scan_quality_dataset(
        dataset,
        now="2026-06-26T12:00:00+08:00",
        vision_analyzer=None,
        pdf_page_renderer=FakePdfPageRenderer(),
    )

    pdf_issues = [issue for issue in scan.issues if issue.file_name == "02496056.pdf"]
    assert calls["pdf_path"] == pdf_path
    assert len(pdf_issues) == 1
    assert pdf_issues[0].rule_id == "R-OCR-001"
    assert pdf_issues[0].preview_image_url == "/api/quality/pdf-pages/02496056-page-1.png"

def test_quality_scan_returns_all_pdf_page_previews_for_manual_review(tmp_path: Path):
    dataset = tmp_path / "pdf-multipage-preview-dataset"
    dataset.mkdir()
    pdf_path = dataset / "report.pdf"
    write_pdf(pdf_path, page_count=3)
    rendered_paths = [tmp_path / f"report-page-{page}.png" for page in range(1, 4)]
    calls: dict[str, Path] = {}

    class FakePdfPageRenderer:
        def render_pages(self, path: Path) -> list[Path]:
            calls["pdf_path"] = path
            for rendered_path in rendered_paths:
                rendered_path.write_bytes(b"\x89PNG\r\n\x1a\nrendered-pdf-page")
            return rendered_paths

    scan = scan_quality_dataset(
        dataset,
        now="2026-06-26T12:00:00+08:00",
        vision_analyzer=None,
        pdf_page_renderer=FakePdfPageRenderer(),
    )

    pdf_issue = [issue for issue in scan.issues if issue.file_name == "report.pdf"][0]
    assert calls["pdf_path"] == pdf_path
    assert pdf_issue.preview_image_url == "/api/quality/pdf-pages/report-page-1.png"
    assert pdf_issue.preview_image_urls == [
        "/api/quality/pdf-pages/report-page-1.png",
        "/api/quality/pdf-pages/report-page-2.png",
        "/api/quality/pdf-pages/report-page-3.png",
    ]
    assert pdf_issue.preview_page_count == 3
def test_quality_scan_uses_rendered_pdf_page_for_vision_model(tmp_path: Path):
    dataset = tmp_path / "pdf-vision-dataset"
    dataset.mkdir()
    pdf_path = dataset / "report.pdf"
    write_pdf(pdf_path, page_count=6)
    rendered_path = tmp_path / "rendered-report-page.png"
    calls: dict[str, Path] = {}

    class FakePdfPageRenderer:
        def render_first_page(self, path: Path) -> Path:
            calls["pdf_path"] = path
            rendered_path.write_bytes(b"\x89PNG\r\n\x1a\nrendered-pdf-page")
            return rendered_path

    class FakeVisionAnalyzer:
        def analyze_image(self, path: Path) -> VisionQualityFinding:
            calls["image_path"] = path
            return VisionQualityFinding(
                issue_type="PDF vision finding",
                severity="medium",
                evidence="Rendered PDF page contains a suspicious abnormal result.",
                ai_judgement="PDF page was analyzed by the vision model.",
                recommendation="Review the original PDF and structured data together.",
                confidence=0.83,
                finding_type="privacy_leak",
                bbox=[8.0, 6.0, 28.0, 10.0],
            )

    scan = scan_quality_dataset(
        dataset,
        now="2026-06-26T12:00:00+08:00",
        vision_analyzer=FakeVisionAnalyzer(),
        pdf_page_renderer=FakePdfPageRenderer(),
    )

    pdf_issues = [issue for issue in scan.issues if issue.file_name == "report.pdf"]
    assert calls["pdf_path"] == pdf_path
    assert calls["image_path"] == rendered_path
    assert len(pdf_issues) == 1
    assert pdf_issues[0].rule_id == "R-VISION-001"
    assert pdf_issues[0].issue_type == "PDF vision finding"
    assert pdf_issues[0].page == "1"
    assert pdf_issues[0].ai_judgement == "PDF page was analyzed by the vision model."
    assert pdf_issues[0].recommendation == "Review the original PDF and structured data together."
    assert pdf_issues[0].category == "privacy"
    assert pdf_issues[0].finding_type == "privacy_leak"
    assert pdf_issues[0].bbox == [8.0, 6.0, 28.0, 10.0]
    assert pdf_issues[0].preview_image_url == "/api/quality/pdf-pages/rendered-report-page.png"




def test_quality_scan_uses_batch_pdf_page_analysis_when_available(tmp_path: Path):
    dataset = tmp_path / "pdf-vision-batch-dataset"
    dataset.mkdir()
    pdf_path = dataset / "report.pdf"
    write_pdf(pdf_path, page_count=6)
    rendered_paths = [tmp_path / f"report-page-{page}.png" for page in range(1, 4)]
    calls: dict[str, object] = {}

    class FakePdfPageRenderer:
        def render_pages(self, path: Path) -> list[Path]:
            for rendered_path in rendered_paths:
                rendered_path.write_bytes(b"\x89PNG\r\n\x1a\nrendered-pdf-page")
            return rendered_paths

    class FakeVisionAnalyzer:
        def analyze_images(self, paths: list[Path]) -> list[tuple[int, VisionQualityFinding]]:
            calls["batch_paths"] = paths
            return [
                (
                    3,
                    VisionQualityFinding(
                        issue_type="Unredacted phone number",
                        severity="high",
                        evidence="Page 3 shows a full phone number.",
                        ai_judgement="The page contains a suspected privacy leak.",
                        recommendation="Mask the phone number before export.",
                        confidence=0.9,
                        finding_type="privacy_leak",
                        bbox=[8.0, 12.0, 38.0, 7.0],
                    ),
                )
            ]

        def analyze_image(self, path: Path) -> VisionQualityFinding:
            raise AssertionError("scan should use batched PDF page analysis")

    scan = scan_quality_dataset(
        dataset,
        now="2026-06-26T12:00:00+08:00",
        vision_analyzer=FakeVisionAnalyzer(),
        pdf_page_renderer=FakePdfPageRenderer(),
    )

    pdf_issues = [issue for issue in scan.issues if issue.file_name == "report.pdf" and issue.rule_id == "R-VISION-001"]
    assert calls["batch_paths"] == rendered_paths
    assert len(pdf_issues) == 1
    assert pdf_issues[0].page == "3"
    assert pdf_issues[0].preview_image_url == "/api/quality/pdf-pages/report-page-3.png"
    assert pdf_issues[0].finding_type == "privacy_leak"


def test_quality_scan_analyzes_all_pdf_pages_and_keeps_only_actionable_findings(tmp_path: Path):
    dataset = tmp_path / "pdf-vision-multipage-dataset"
    dataset.mkdir()
    pdf_path = dataset / "report.pdf"
    write_pdf(pdf_path, page_count=6)
    rendered_paths = [tmp_path / f"report-page-{page}.png" for page in range(1, 4)]
    analyzed_paths: list[Path] = []

    class FakePdfPageRenderer:
        def render_pages(self, path: Path) -> list[Path]:
            assert path == pdf_path
            for rendered_path in rendered_paths:
                rendered_path.write_bytes(b"\x89PNG\r\n\x1a\nrendered-pdf-page")
            return rendered_paths

    class FakeVisionAnalyzer:
        def analyze_image(self, path: Path) -> VisionQualityFinding:
            analyzed_paths.append(path)
            if path == rendered_paths[0]:
                return VisionQualityFinding(
                    issue_type="No actionable finding",
                    severity="low",
                    evidence="The first page is blank or has no visible report result.",
                    ai_judgement="No content-level issue can be confirmed from this page.",
                    recommendation="Continue reviewing the remaining pages.",
                    confidence=0.2,
                    finding_type="other",
                    bbox=None,
                    has_finding=False,
                )
            if path == rendered_paths[1]:
                return VisionQualityFinding(
                    issue_type="Abnormal glucose value needs review",
                    severity="high",
                    evidence="Page 2 shows glucose 12.8 mmol/L above the reference range.",
                    ai_judgement="The visible value is a content-level abnormal result that should be reviewed.",
                    recommendation="Review page 2 and compare the structured glucose value.",
                    confidence=0.92,
                    finding_type="data_error",
                    bbox=[16.0, 42.0, 28.0, 8.0],
                )
            return VisionQualityFinding(
                issue_type="Unredacted identity number",
                severity="high",
                evidence="Page 3 shows a full identity number near the patient information area.",
                ai_judgement="The page contains a suspected privacy leak.",
                recommendation="Mask the identity number before export or sharing.",
                confidence=0.9,
                finding_type="privacy_leak",
                bbox=[8.0, 12.0, 38.0, 7.0],
            )

    scan = scan_quality_dataset(
        dataset,
        now="2026-06-26T12:00:00+08:00",
        vision_analyzer=FakeVisionAnalyzer(),
        pdf_page_renderer=FakePdfPageRenderer(),
    )

    pdf_issues = [issue for issue in scan.issues if issue.file_name == "report.pdf" and issue.rule_id == "R-VISION-001"]
    assert analyzed_paths == rendered_paths
    assert [issue.page for issue in pdf_issues] == ["2", "3"]
    assert [issue.issue_type for issue in pdf_issues] == ["Abnormal glucose value needs review", "Unredacted identity number"]
    assert pdf_issues[0].bbox == [16.0, 42.0, 28.0, 8.0]
    assert pdf_issues[1].finding_type == "privacy_leak"
    assert all(issue.preview_image_urls == [
        "/api/quality/pdf-pages/report-page-1.png",
        "/api/quality/pdf-pages/report-page-2.png",
        "/api/quality/pdf-pages/report-page-3.png",
    ] for issue in pdf_issues)
    assert all(issue.preview_page_count == 3 for issue in pdf_issues)

def test_quality_scan_maps_target_visual_quality_findings_to_review_categories(tmp_path: Path):
    dataset = tmp_path / "pdf-target-findings-dataset"
    dataset.mkdir()
    pdf_path = dataset / "report.pdf"
    write_pdf(pdf_path, page_count=6)
    rendered_paths = [tmp_path / f"target-page-{page}.png" for page in range(1, 5)]

    class FakePdfPageRenderer:
        def render_pages(self, path: Path) -> list[Path]:
            assert path == pdf_path
            for rendered_path in rendered_paths:
                rendered_path.write_bytes(b"\x89PNG\r\n\x1a\nrendered-pdf-page")
            return rendered_paths

    class FakeVisionAnalyzer:
        def analyze_images(self, paths: list[Path]) -> list[tuple[int, VisionQualityFinding]]:
            assert paths == rendered_paths
            return [
                (
                    1,
                    VisionQualityFinding(
                        issue_type="疑似缺字",
                        severity="high",
                        evidence="第 1 页检验结果栏有关键字缺损，无法完整确认项目名称。",
                        ai_judgement="页面存在疑似缺字，需要人工复核原图。",
                        recommendation="查看框选区域，回查原 PDF 或重新扫描件。",
                        confidence=0.9,
                        finding_type="missing_text",
                        bbox=[12.0, 18.0, 30.0, 8.0],
                    ),
                ),
                (
                    2,
                    VisionQualityFinding(
                        issue_type="疑似未脱敏",
                        severity="high",
                        evidence="第 2 页顶部可见完整身份证号。",
                        ai_judgement="页面存在疑似未脱敏个人身份信息。",
                        recommendation="进入人工复核，不自动修改 PDF。",
                        confidence=0.92,
                        finding_type="privacy_leak",
                        bbox=[8.0, 6.0, 38.0, 7.0],
                    ),
                ),
                (
                    3,
                    VisionQualityFinding(
                        issue_type="历史对比缺失",
                        severity="medium",
                        evidence="第 3 页写有历史对比，但未展示历史指标或对比表。",
                        ai_judgement="历史对比证据缺失，需要人工确认是否满足客户要求。",
                        recommendation="回查同一档案号历史报告和结构化数据。",
                        confidence=0.86,
                        finding_type="history_gap",
                        bbox=[18.0, 48.0, 45.0, 10.0],
                    ),
                ),
                (
                    4,
                    VisionQualityFinding(
                        issue_type="页面边界裁切",
                        severity="medium",
                        evidence="第 4 页底部结论区域被裁切。",
                        ai_judgement="页面格式存在边界裁切风险。",
                        recommendation="复核扫描页边界并确认是否缺页。",
                        confidence=0.84,
                        finding_type="format",
                        bbox=[5.0, 88.0, 90.0, 9.0],
                    ),
                ),
            ]

    scan = scan_quality_dataset(
        dataset,
        now="2026-06-26T12:00:00+08:00",
        vision_analyzer=FakeVisionAnalyzer(),
        pdf_page_renderer=FakePdfPageRenderer(),
    )

    findings = [issue for issue in scan.issues if issue.rule_id == "R-VISION-001"]
    assert [issue.issue_type for issue in findings] == ["疑似缺字", "疑似未脱敏", "历史对比缺失", "页面边界裁切"]
    assert [issue.category for issue in findings] == ["content", "privacy", "history", "format"]
    assert [issue.finding_type for issue in findings] == ["missing_text", "privacy_leak", "history_gap", "format"]
    assert all(issue.status == "needs_review" for issue in findings)
    assert all(issue.preview_image_urls for issue in findings)
def test_quality_api_imports_images_and_serves_preview(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "local_storage_root", str(tmp_path / "api-state"))
    client = TestClient(app)

    response = client.post(
        "/api/quality/import",
        files=[
            (
                "files",
                (
                    "report-page.png",
                    b"\x89PNG\r\n\x1a\nimage-bytes",
                    "image/png",
                ),
            )
        ],
    )

    assert response.status_code == 201
    body = response.json()

    scan_response = client.post(f"/api/quality/imports/{body['task_id']}/scan")

    assert scan_response.status_code == 200
    scan_body = scan_response.json()
    image_issues = [
        issue
        for issue in scan_body["scan"]["issues"]
        if issue["file_name"] == "report-page.png"
    ]
    assert image_issues
    assert image_issues[0]["rule_id"] == "R-IMAGE-001"
    assert image_issues[0]["issue_type"] == "图片待 OCR"

    preview_response = client.get(f"/api/quality/imports/{body['task_id']}/files/report-page.png")

    assert preview_response.status_code == 200
    assert preview_response.headers["content-type"].startswith("image/png")
    assert preview_response.content == b"\x89PNG\r\n\x1a\nimage-bytes"
def test_quality_api_serves_pdf_page_preview(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "local_storage_root", str(tmp_path / "api-state"))
    preview_dir = Path(settings.local_storage_root) / "quality" / "pdf-pages"
    preview_dir.mkdir(parents=True)
    (preview_dir / "rendered-report-page.png").write_bytes(b"\x89PNG\r\n\x1a\nrendered-pdf-page")
    client = TestClient(app)

    response = client.get("/api/quality/pdf-pages/rendered-report-page.png")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("image/png")
    assert response.content == b"\x89PNG\r\n\x1a\nrendered-pdf-page"

def test_quality_import_scan_response_preserves_page_level_evidence_fields(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "local_storage_root", str(tmp_path / "api-state"))
    task_id = "quality-import-evidence"
    dataset_dir = Path(settings.local_storage_root) / "quality" / "imports" / task_id
    dataset_dir.mkdir(parents=True)
    (dataset_dir / "report-page.png").write_bytes(b"\x89PNG\r\n\x1a\nimage-bytes")

    def fake_scan_quality_dataset(dataset_path: Path) -> ServiceQualityDatasetScan:
        return ServiceQualityDatasetScan(
            dataset_path=str(dataset_path),
            scanned_at="2026-06-26T12:00:00+08:00",
            metrics=[
                ServiceQualityMetric("??", "1", "warning", "orange"),
            ],
            pipeline=[
                ServiceQualityPipelineStage("????", "1/1", True, False),
            ],
            issues=[
                ServiceQualityIssue(
                    id="r-vision-001-1-report-page",
                    index=1,
                    file_name="report-page.png",
                    group="",
                    archive_id="report-page",
                    page="??",
                    issue_type="?????????",
                    category="content",
                    severity="high",
                    rule_id="R-VISION-001",
                    evidence="?????????????????",
                    ai_judgement="?????????????????",
                    recommendation="???????????????",
                    confidence=0.91,
                    status="needs_review",
                    found_at="2026-06-26 12:00:00",
                    finding_type="data_error",
                    bbox=[12.0, 18.0, 36.0, 14.0],
                    preview_image_url="/api/quality/imports/quality-import-evidence/files/report-page.png",
                )
            ],
        )

    monkeypatch.setattr(quality_router, "scan_quality_dataset", fake_scan_quality_dataset)
    client = TestClient(app)

    response = client.post(f"/api/quality/imports/{task_id}/scan")

    assert response.status_code == 200
    issue = response.json()["scan"]["issues"][0]
    assert issue["finding_type"] == "data_error"
    assert issue["bbox"] == [12.0, 18.0, 36.0, 14.0]
    assert issue["preview_image_url"] == "/api/quality/imports/quality-import-evidence/files/report-page.png"


def test_quality_review_and_export_records_are_persisted(tmp_path: Path):
    dataset = build_quality_sample(tmp_path)
    storage_root = tmp_path / "quality-state"
    review = record_review_decision(
        issue_id="issue-001",
        decision="confirmed",
        reviewer="backend-tester",
        note="confirmed",
        evidence="PDF page count is 3",
        dataset_path=dataset,
        storage_root=storage_root,
        now="2026-06-26T12:30:00+08:00",
    )
    export = create_export_task(
        "review-records",
        dataset_path=dataset,
        storage_root=storage_root,
        now="2026-06-26T12:31:00+08:00",
    )

    assert list_review_records(storage_root=storage_root) == [review]
    assert list_review_records(dataset_path=dataset, storage_root=storage_root) == [review]
    assert list_export_tasks(storage_root=storage_root) == [export]
    assert (storage_root / "quality" / "review-records.jsonl").exists()
    assert (storage_root / "quality" / "export-tasks.jsonl").exists()
    assert Path(export.bundle_path or "").exists()


def test_quality_export_task_generates_real_delivery_package(tmp_path: Path):
    dataset = build_quality_sample(tmp_path)
    storage_root = tmp_path / "quality-state"
    scan = scan_quality_dataset(dataset, now="2026-06-30T12:00:00+08:00")

    record_review_decision(
        issue_id=scan.issues[0].id,
        decision="confirmed",
        reviewer="reviewer-a",
        note="confirmed",
        evidence=scan.issues[0].evidence,
        dataset_path=dataset,
        storage_root=storage_root,
        now="2026-06-30T12:10:00+08:00",
    )
    record_review_decision(
        issue_id=scan.issues[1].id,
        decision="rejected",
        reviewer="reviewer-b",
        note="rejected",
        evidence=scan.issues[1].evidence,
        dataset_path=dataset,
        storage_root=storage_root,
        now="2026-06-30T12:11:00+08:00",
    )

    export = create_export_task(
        "第三批数据检测报告",
        dataset_path=dataset,
        storage_root=storage_root,
        now="2026-06-30T12:12:00+08:00",
    )

    assert export.status == "done"
    assert export.artifact_count and export.artifact_count >= 7
    assert export.bundle_name and export.bundle_name.endswith(".zip")
    assert export.bundle_path
    assert export.download_url == f"/api/quality/exports/{export.id}/download"

    bundle_path = Path(export.bundle_path)
    assert bundle_path.exists()

    with ZipFile(bundle_path) as archive:
        names = set(archive.namelist())
        assert any(name.endswith("third-batch-report.md") for name in names)
        assert any(name.endswith("non-compliant-issues.jsonl") for name in names)
        assert any(name.endswith("possible-compliant-issues.jsonl") for name in names)
        assert any(name.endswith("needs-review-issues.jsonl") for name in names)
        assert any(name.endswith("review-records.jsonl") for name in names)
        assert any(name.endswith("rule-hit-stats.json") for name in names)
        assert any(name.endswith("evidence-image-index.json") for name in names)
        assert any(name.endswith("export-summary.json") for name in names)
        report_path = next(name for name in names if name.endswith("third-batch-report.md"))
        report = archive.read(report_path).decode("utf-8")
        assert "第三批数据检测报告" in report
        assert "总问题：3" in report
        confirmed_path = next(name for name in names if name.endswith("non-compliant-issues.jsonl"))
        rejected_path = next(name for name in names if name.endswith("possible-compliant-issues.jsonl"))
        pending_path = next(name for name in names if name.endswith("needs-review-issues.jsonl"))
        confirmed = archive.read(confirmed_path).decode("utf-8")
        rejected = archive.read(rejected_path).decode("utf-8")
        pending = archive.read(pending_path).decode("utf-8")
        assert "R-FORMAT-002" in confirmed
        assert "R-OCR-001" in rejected
        assert "R-EXCEL-001" in pending


def test_quality_api_can_download_generated_export_bundle(tmp_path: Path, monkeypatch):
    dataset = build_quality_sample(tmp_path)
    storage_root = tmp_path / "api-state"
    monkeypatch.setattr(settings, "local_storage_root", str(storage_root))
    client = TestClient(app)

    create_response = client.post(
        "/api/quality/exports",
        json={"export_type": "第三批数据检测报告", "dataset_path": str(dataset)},
    )

    assert create_response.status_code == 200
    export_id = create_response.json()["id"]

    download_response = client.get(f"/api/quality/exports/{export_id}/download")

    assert download_response.status_code == 200
    assert download_response.headers["content-type"].startswith("application/zip")
    assert len(download_response.content) > 0

def test_quality_export_summary_aggregates_real_assets_issues_and_reviews(tmp_path: Path):
    dataset = build_quality_sample(tmp_path)
    storage_root = tmp_path / "quality-state"
    scan = scan_quality_dataset(dataset, now="2026-06-30T12:00:00+08:00")

    record_review_decision(
        issue_id=scan.issues[0].id,
        decision="confirmed",
        reviewer="reviewer-a",
        note="confirmed",
        evidence=scan.issues[0].evidence,
        dataset_path=dataset,
        storage_root=storage_root,
        now="2026-06-30T12:10:00+08:00",
    )
    record_review_decision(
        issue_id=scan.issues[1].id,
        decision="rejected",
        reviewer="reviewer-b",
        note="rejected",
        evidence=scan.issues[1].evidence,
        dataset_path=dataset,
        storage_root=storage_root,
        now="2026-06-30T12:11:00+08:00",
    )

    summary = build_quality_export_summary(
        dataset,
        storage_root=storage_root,
        now="2026-06-30T12:12:00+08:00",
    )

    assert isinstance(summary, ServiceQualityExportSummary)
    assert summary.dataset_path == str(dataset)
    assert summary.total_issues == 3
    assert summary.confirmed_issues == 1
    assert summary.rejected_issues == 1
    assert summary.pending_issues == 1
    assert summary.review_record_count == 2
    assert summary.evidence_image_count == 3
    assert [section.key for section in summary.sections] == [
        "third-batch-report",
        "non-compliant",
        "possible-compliant",
        "needs-review",
        "review-records",
        "rule-hit-stats",
        "batch-overview-table",
        "structured-data",
        "compliant-pdfs",
        "issue-detail-reports",
    ]
    assert [section.item_count for section in summary.sections[1:5]] == [1, 1, 1, 2]
    assert [rule_hit.rule_id for rule_hit in summary.rule_hits] == [
        "R-EXCEL-001",
        "R-FORMAT-002",
        "R-OCR-001",
    ]
    assert all(rule_hit.hit_count == 1 for rule_hit in summary.rule_hits)


def test_quality_api_exposes_export_summary(tmp_path: Path, monkeypatch):
    dataset = build_quality_sample(tmp_path)
    storage_root = tmp_path / "api-state"
    monkeypatch.setattr(settings, "local_storage_root", str(storage_root))
    scan = scan_quality_dataset(dataset, now="2026-06-30T12:00:00+08:00")
    record_review_decision(
        issue_id=scan.issues[0].id,
        decision="confirmed",
        reviewer="api-reviewer",
        note="confirmed",
        evidence=scan.issues[0].evidence,
        dataset_path=dataset,
        storage_root=storage_root,
        now="2026-06-30T12:05:00+08:00",
    )
    client = TestClient(app)

    response = client.get("/api/quality/exports/summary", params={"dataset_path": str(dataset)})

    assert response.status_code == 200
    body = response.json()
    assert body["dataset_path"] == str(dataset)
    assert body["total_issues"] == 3
    assert body["confirmed_issues"] == 1
    assert body["review_record_count"] == 1
    assert {section["key"] for section in body["sections"]} == {
        "third-batch-report",
        "non-compliant",
        "possible-compliant",
        "needs-review",
        "review-records",
        "rule-hit-stats",
        "batch-overview-table",
        "structured-data",
        "compliant-pdfs",
        "issue-detail-reports",
    }


def test_quality_export_does_not_render_pdf_pages_during_package_generation(tmp_path: Path, monkeypatch):
    dataset = build_quality_sample(tmp_path)

    def fail_if_renderer_is_created():
        raise AssertionError("export package generation must not render PDF pages")

    monkeypatch.setattr("store_ai_clinic.services.quality._default_pdf_page_renderer", fail_if_renderer_is_created)

    export = create_export_task(
        "快速交付包",
        dataset_path=dataset,
        storage_root=tmp_path / "fast-export-state",
        selected_sections=["issue-detail-reports", "export-summary"],
        now="2026-06-30T12:20:00+08:00",
    )

    assert export.status == "done"
    assert export.bundle_path

def test_quality_export_creates_customer_readable_html_entrypoints(tmp_path: Path):
    dataset = build_quality_sample(tmp_path)
    storage_root = tmp_path / "customer-export-state"

    export = create_export_task(
        "客户交付包",
        dataset_path=dataset,
        storage_root=storage_root,
        selected_sections=[
            "third-batch-report",
            "batch-overview-table",
            "issue-detail-reports",
            "export-summary",
        ],
        now="2026-06-30T12:20:00+08:00",
    )

    assert export.bundle_name is not None
    assert export.bundle_name.startswith("体检报告质检交付包_sample-dataset_20260630_1220")
    assert export.bundle_path is not None

    with ZipFile(Path(export.bundle_path)) as archive:
        names = set(archive.namelist())
        assert "00-交付包说明.html" in names
        assert "01-批次质检总报告.html" in names
        assert any(name.startswith("02-逐份报告问题说明/") and name.endswith(".html") for name in names)

        index_html = archive.read("00-交付包说明.html").decode("utf-8")
        assert "sample-dataset" in index_html
        assert "客户阅读顺序" in index_html
        assert "01-批次质检总报告.html" in index_html
        assert "02-逐份报告问题说明" in index_html

        report_html = archive.read("01-批次质检总报告.html").decode("utf-8")
        assert "批次质检总报告" in report_html
        assert "02496166_report.pdf" in report_html
        assert "问题分类统计" in report_html

        detail_htmls = [
            archive.read(name).decode("utf-8")
            for name in names
            if name.startswith("02-逐份报告问题说明/") and name.endswith(".html")
        ]
        pdf_detail_html = next(content for content in detail_htmls if "02496166_report.pdf" in content)
        assert "命中规则" in pdf_detail_html
        assert "处理建议" in pdf_detail_html

def test_quality_export_batch_deliverables(tmp_path: Path):
    dataset = build_quality_sample(tmp_path)
    storage_root = tmp_path / "batch-export-state"
    export = create_export_task(
        "批量检测交付包",
        dataset_path=dataset,
        storage_root=storage_root,
        selected_sections=[
            "batch-overview-table",
            "structured-data",
            "compliant-pdfs",
            "issue-detail-reports",
            "export-summary",
        ],
        now="2026-06-30T12:20:00+08:00",
    )

    assert export.status == "done"
    assert export.bundle_path
    assert export.export_dir
    assert Path(export.export_dir).exists()
    assert (Path(export.export_dir) / "批次总体情况表").is_dir()
    assert (Path(export.export_dir) / "结构化数据").is_dir()

    with ZipFile(Path(export.bundle_path)) as archive:
        names = set(archive.namelist())
        assert any("批次总体情况表" in name and name.endswith(".xlsx") for name in names)
        assert any("结构化数据" in name and name.endswith(".xlsx") for name in names)
        assert any("问题详情分析" in name and name.endswith(".md") for name in names)
        overview_path = next(name for name in names if "批次总体情况表" in name and name.endswith(".xlsx"))
        from openpyxl import load_workbook
        from io import BytesIO

        workbook = load_workbook(BytesIO(archive.read(overview_path)))
        sheet = workbook.active
        assert sheet.cell(1, 1).value == "年龄段"
        assert sheet.cell(1, 8).value == "合规状态"


